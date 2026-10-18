"""Build the editable EchoSphere pitch deck (python-pptx; no runtime app changes)."""
from pathlib import Path
import math
import random
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
BG = '090D1B'
PANEL = '11182A'
BORDER = '29324C'
WHITE = 'F4F2FF'
MUTED = 'B2B9D0'
PURPLE = 'BCABFF'
CYAN = '86DAF2'
ORANGE = 'FF967B'
GREEN = 'B5E58C'
FONT = 'Arial'
WIDTH, HEIGHT = 13.333333, 7.5

CONTENT = [
    {
        'title': 'EchoSphere', 'seconds': 25, 'window': '0:00–0:25',
        'script': "EchoSphere explores a simple idea: your voice should do more than choose the next line of a story. It should shape the world around you. We combine continuous audio feedback with a narrative guide, using an Evolus agent connected to Gemma on AMD Developer Cloud.",
        'cue': 'Open on the title. Emphasize the three technology names; no live app switch is needed.',
    },
    {
        'title': 'Who it is for', 'seconds': 30, 'window': '0:25–0:55',
        'script': "We are building for immersive studios, museums, and creators of interactive audio. They create the worlds; visitors and listeners explore them. The aim is to make one story world respond differently to each participant, while offering voice, text, and keyboard controls. Our proposed business model is a creator or venue subscription. That is a customer hypothesis we still need to validate.",
        'cue': 'Point from the two creator/customer cards to the participant card.',
    },
    {
        'title': 'The problem it addresses', 'seconds': 35, 'window': '0:55–1:30',
        'script': "A transcript captures what we say, but loses much of how we say it. The same words can arrive softly, loudly, or with long pauses. A story driven only by text has no direct signal for those differences. And if every reaction waits for a completed turn, the world can feel unresponsive while we speak. EchoSphere separates immediate acoustic feedback from slower narrative reasoning.",
        'cue': 'Use the two waveforms to illustrate delivery only. Do not assign emotions to either waveform.',
    },
    {
        'title': 'What it does', 'seconds': 45, 'window': '1:30–2:15',
        'script': "The experience opens in a starfield around a draggable, sixteen-orbit atom. While you speak, local audio frames change the visuals and ambience. A slower, parallel path recognizes your words and sends them, with acoustic context, to the narrative agent. The guide leads the story, and you can interrupt it. Browser-generated speech passes through spatial positioning, distance filtering, and room effects. Story cues choose the placement, while vocal delivery modulates it. Words move the plot; the way you speak shapes the atmosphere.",
        'cue': 'The screenshot is the real app interface. Describe the implemented paths; do not imply that a live speech-model acceptance test has passed.',
    },
    {
        'title': 'The role of technology', 'seconds': 55, 'window': '2:15–3:10',
        'script': "Each technology has a distinct role. AMD Developer Cloud hosts the meaningful compute workload: narrative inference. The configured Gemma open model generates the story language using participant input and recent context. Evolus provides the agent and API connection around that model. Our Node backend keeps the credential on the server and forwards streamed responses. In the browser, Whisper handles recognition, Kokoro generates speech, and Web Audio positions that speech in space. The detector, interface, and placement rules are our authored work; narrative prose is model generated. Render hosts the web app, while the narrative model remains on AMD.",
        'cue': 'Show the actual AMD deployment/model evidence separately with secrets hidden if available. Do not name a GPU SKU, ROCm version, or serving stack that has not been verified.',
    },
    {
        'title': 'Two inputs. A different world.', 'seconds': 60, 'window': '3:10–4:10',
        'script': "We have tested two written inputs against the real Evolus agent. A quiet request produced a guide right beside the listener; a distant encounter produced a reply from the far right. The local director mapped those replies to different spatial presets. These are real text responses, not proof of microphone recognition. Fifty-four automated tests pass, including audio and backend checks. Our remaining acceptance steps are successful real Whisper and Kokoro inference, followed by the public HTTPS voice demo. Costs come from AMD compute, hosting, and any billable Evolus services; cost per session is not yet measured. EchoSphere's goal is simple: turn listening into participation.",
        'cue': 'These examples are from the documented live typed comparison. Keep the verification line visible. Finish by 4:10; the remaining 20 seconds are a buffer, not another demo segment.',
    },
]

prs = Presentation()
prs.slide_width = Inches(WIDTH)
prs.slide_height = Inches(HEIGHT)
prs.core_properties.title = 'EchoSphere — Your voice shapes the story'
prs.core_properties.subject = 'AMD Developer Hackathon Act 3 · Track 4'
prs.core_properties.author = 'EchoSphere'
prs.core_properties.keywords = 'EchoSphere, AMD Developer Cloud, Gemma, Evolus, spatial audio'


def rgb(c):
    return RGBColor.from_string(c)


def blend(c, alpha, base=BG):
    a = [int(c[i:i+2], 16) for i in (0, 2, 4)]
    b = [int(base[i:i+2], 16) for i in (0, 2, 4)]
    return ''.join(f'{round(x*alpha+y*(1-alpha)):02X}' for x, y in zip(a, b))


def rect(s, x, y, w, h, fill=PANEL, line=None, radius=False):
    q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                           Inches(x), Inches(y), Inches(w), Inches(h))
    q.fill.solid(); q.fill.fore_color.rgb = rgb(fill)
    q.line.fill.background() if line is None else None
    if line:
        q.line.color.rgb = rgb(line); q.line.width = Pt(.8)
    if radius:
        q.adjustments[0] = .09
    return q


def line(s, x1, y1, x2, y2, color=BORDER, width=1):
    q = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    q.line.color.rgb = rgb(color); q.line.width = Pt(width)
    return q


def circle(s, x, y, r, color):
    q = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r), Inches(2*r), Inches(2*r))
    q.fill.solid(); q.fill.fore_color.rgb = rgb(color); q.line.fill.background()
    return q


def text(s, value, x, y, w, h, size=20, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, row in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = row; p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(3)
        p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = rgb(color)
    return box


def path(s, points, color, width=1):
    # Native vector path, with editable stroke; inch coordinates -> drawing units.
    f = s.shapes.build_freeform(round(points[0][0]*1000), round(points[0][1]*1000), Inches(.001))
    f.add_line_segments([(round(x*1000), round(y*1000)) for x, y in points[1:]], close=False)
    q = f.convert_to_shape(); q.fill.background()
    q.line.color.rgb = rgb(color); q.line.width = Pt(width)
    return q


def waveform(s, x, y, w, amplitude, color, phase=0):
    points = []
    for i in range(161):
        t = i/160
        envelope = math.sin(math.pi*t)**1.6
        v = (math.sin(t*7*math.pi+phase) + .34*math.sin(t*15*math.pi+phase)) / 1.34
        points.append((x+t*w, y+v*amplitude*envelope))
    line(s, x, y, x+w, y, blend(color, .22), .65)
    path(s, points, color, 1.7)


def atom(s, cx, cy, radius):
    # Perspective projection of sixteen orbital planes, matching the app's motif.
    for shell in range(16):
        a = shell*math.pi/16
        tilt = .3 + .65*abs(math.sin(a*2+.3))
        rot = a
        r = radius*(.79+.2*math.sin(shell*2.4)**2)
        def project(t):
            u, v = math.cos(t)*r, math.sin(t)*r*tilt
            return (cx+u*math.cos(rot)-v*math.sin(rot), cy+u*math.sin(rot)+v*math.cos(rot))
        path(s, [project(i*math.tau/100) for i in range(101)], blend(PURPLE if shell%2 else CYAN, .24), .7)
        count = 14 if shell < 14 else 13
        for i in range(count):
            ex, ey = project(i*math.tau/count+shell*.16)
            circle(s, ex, ey, .014 + .008*((i+shell)%3==0), blend(CYAN if shell%3 else PURPLE, .45+.4*((i+shell)%3==0)))
    for layer in range(5, 0, -1):
        circle(s, cx, cy, layer*.08, blend(PURPLE, .035*layer))
    for i in range(13):
        r = .16*math.sqrt(i/13)
        circle(s, cx+r*math.cos(i*2.4), cy+r*math.sin(i*2.4), .052, [PURPLE,CYAN,'E9D9FF'][i%3])


def base(n, eyebrow):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    rng = random.Random(70+n)
    for _ in range(65):
        x,y = rng.uniform(.15,13.15),rng.uniform(.15,7.25)
        circle(s,x,y,rng.choice([.007,.009,.012]),blend(PURPLE,rng.uniform(.12,.38)))
    text(s, 'E C H O S P H E R E', .58, .35, 3.7, .22, 10, WHITE, True)
    text(s, eyebrow, 6.2, .35, 6.55, .22, 10, MUTED, False, PP_ALIGN.RIGHT)
    line(s, .58, .79, 12.75, .79, BORDER, .65)
    for i in range(6):
        rect(s, .58+i*.31, 7.12, .22, .025, PURPLE if i==n-1 else BORDER)
    text(s, 'AMD DEVELOPER HACKATHON ACT 3  /  TRACK 4', 3, 7.01, 7.2, .23, 9, MUTED)
    text(s, f'{n:02} / 06', 11.85, 7.01, .9, .23, 10, MUTED, False, PP_ALIGN.RIGHT)
    c = CONTENT[n-1]
    s.notes_slide.notes_text_frame.text = f"{c['window']} · {c['seconds']} seconds\n\nSAY\n{c['script']}\n\nPRESENTER CUE (not spoken)\n{c['cue']}\n\nTarget total: 4:10. Hard limit: 4:30."
    transition=OxmlElement('p:transition')
    transition.set('advClick','1'); transition.set('advTm',str(c['seconds']*1000))
    s._element.insert_element_before(transition,'p:timing','p:extLst')
    return s


# 01 / Title. Technology appears immediately, not only on the architecture slide.
s=base(1,'CONTINUOUS AUDIO · SPATIAL STORYTELLING')
text(s,'EchoSphere',.65,1.38,7.6,1.05,62,WHITE,True)
text(s,'Your voice shapes\nthe story’s atmosphere.',.69,2.64,6.8,1.35,31,PURPLE)
text(s,'An audio-first narrative experiment.',.72,4.27,6,.4,17,MUTED)
atom(s,10.0,3.55,2.22)
text(s,'VOICE  →  ATMOSPHERE  →  STORY',7.55,6.05,4.7,.3,11,MUTED,False,PP_ALIGN.CENTER)
for x,w,name,sub,col in [(.7,2.7,'AMD','Developer Cloud',ORANGE),(3.62,2.05,'Gemma','Open model',CYAN),(5.89,1.95,'Evolus','Agent + API',GREEN)]:
    line(s,x,5.42,x+w,5.42,col,1.8)
    text(s,name,x,5.61,w,.43,25,col,True)
    text(s,sub,x,6.1,w,.28,12,MUTED)

# 02 / Audience. Customers and participants are deliberately distinguished.
s=base(2,'WHO IT IS FOR')
text(s,'For the people who\nbring stories to life.',.65,1.12,11.6,1.35,40,WHITE,True)
cards=[(.68,'01','Immersive studios','Create responsive worlds\nfor interactive entertainment.',PURPLE),
       (4.77,'02','Museums & venues','Turn a visit into a story\nthe audience can influence.',CYAN),
       (8.86,'03','Audio-first audiences','Explore through voice, with\ntext and keyboard alternatives.',GREEN)]
for x,num,title,body,c in cards:
    rect(s,x,3.05,3.81,2.47,PANEL,BORDER,True)
    text(s,num,x+.24,3.29,.6,.4,21,c,True)
    text(s,title,x+.24,4.0,3.33,.4,21,WHITE,True)
    text(s,body,x+.24,4.62,3.33,.65,15,MUTED)
text(s,'CREATORS / POTENTIAL CUSTOMERS',.72,5.74,7.7,.26,10,PURPLE,True)
text(s,'PARTICIPANTS',8.9,5.74,3.5,.26,10,GREEN,True)
text(s,'Proposed model: creator or venue subscription.',.7,6.38,11.5,.35,17,MUTED)

# 03 / Problem. Same semantic content, different delivery; no emotion labels.
s=base(3,'THE PROBLEM IT ADDRESSES')
text(s,'A voice carries\nmore than words.',.65,1.17,6.5,1.4,42,WHITE,True)
text(s,'A transcript loses softness,\npauses and intensity.',.7,3.12,5.6,.92,24,MUTED)
line(s,.7,4.47,5.55,4.47,PURPLE,1)
text(s,'Waiting for a full turn can leave\nthe world unresponsive while we speak.',.7,4.76,5.65,1.0,21,WHITE)
rect(s,7.06,1.36,5.61,4.84,PANEL,BORDER,True)
text(s,'SAME WORDS',7.4,1.72,4.9,.26,11,MUTED,True)
text(s,'“I’m here.”',7.4,2.15,4.9,.65,35,WHITE,True)
waveform(s,7.42,3.42,4.87,.15,CYAN)
text(s,'SOFT DELIVERY',7.43,3.75,4.85,.3,11,CYAN,True)
waveform(s,7.42,4.8,4.87,.39,PURPLE,.3)
text(s,'STRONGER DELIVERY',7.43,5.42,4.85,.3,11,PURPLE,True)
text(s,'Delivery is an acoustic cue, not a diagnosis of emotion.',7.08,6.38,5.61,.35,12,MUTED)

# 04 / Product. Real product capture plus the two parallel response paths.
s=base(4,'WHAT IT DOES')
text(s,'Words move the plot.\nYour delivery shapes the world.',.65,1.07,12,1.28,37,WHITE,True)
rect(s,.67,2.72,5.53,3.85,PANEL,BORDER,True)
s.shapes.add_picture(str(ROOT/'docs/preview.png'),Inches(.81),Inches(2.86),width=Inches(5.25),height=Inches(3.646))
text(s,'01',6.69,2.88,.55,.4,23,CYAN,True)
text(s,'As you speak',7.34,2.86,5.1,.42,25,CYAN,True)
text(s,'50 ms audio frames shape the atom,\nwaveform and ambient sound locally.',7.34,3.46,5.1,.8,19,WHITE)
line(s,6.7,4.43,12.65,4.43,BORDER)
text(s,'02',6.69,4.74,.55,.4,23,PURPLE,True)
text(s,'As the story develops',7.34,4.72,5.1,.42,25,PURPLE,True)
text(s,'A guide leads, responds and can be interrupted.\nStory cues direct a spatial voice.',7.34,5.32,5.1,.83,18,WHITE)
text(s,'16 orbits · Voice + text input · Automatic scene placement',6.7,6.42,6.0,.28,12,MUTED)

# 05 / Technology. Brand text is not a claim of endorsement.
s=base(5,'THE ROLE OF TECHNOLOGY')
text(s,'AMD runs the narrative.\nThe browser shapes the moment.',.65,1.07,12,1.28,37,WHITE,True)
for x,w,title,sub,body,col in [(.69,3.82,'AMD','DEVELOPER CLOUD','GPU compute for\nnarrative inference.',ORANGE),
                            (4.76,3.82,'Gemma','CONFIGURED: GEMMA 4 · 31B IT','Generates story language\nfrom input and context.',CYAN),
                            (8.83,3.82,'Evolus','AGENT + API','Connects the app to\nthe configured model.',GREEN)]:
    rect(s,x,2.82,w,2.34,PANEL,BORDER,True)
    rect(s,x+.24,3.08,.045,1.77,col)
    text(s,title,x+.43,3.01,w-.65,.55,31,col,True)
    text(s,sub,x+.45,3.7,w-.72,.25,10,MUTED,True)
    text(s,body,x+.45,4.22,w-.7,.65,17,WHITE)
text(s,'REQUEST PATH',.73,5.58,2,.25,10,MUTED,True)
text(s,'Browser  →  Secure backend  →  Evolus  →  Gemma on AMD',2.8,5.51,9.9,.4,20,WHITE,True)
text(s,'LOCAL AUDIO',.73,6.22,2,.25,10,MUTED,True)
text(s,'Whisper: words in     ·     Kokoro: voice out     ·     Web Audio: spatial sound',2.8,6.19,9.9,.4,15,PURPLE)

# 06 / Measured evidence and explicit readiness boundary.
s=base(6,'TWO INPUTS · TWO STORY RESPONSES')
text(s,'Turn listening into participation.',.65,1.16,12.0,.85,39,WHITE,True)
for x,tag,quote,result,col in [(.7,'QUIET DISCOVERY','“Please stay beside me\nand speak softly.”','Guide close by\n0.65 m placement preset',CYAN),
                              (6.89,'DISTANT ENCOUNTER','“Let a distant voice answer\nfrom the right.”','Voice from the far right\n7 m / 55° placement preset',PURPLE)]:
    rect(s,x,2.46,5.74,2.71,PANEL,BORDER,True)
    text(s,tag,x+.27,2.74,5.2,.26,11,col,True)
    text(s,quote,x+.27,3.25,5.2,.8,22,WHITE,True)
    text(s,result,x+.27,4.34,5.2,.65,16,col)
text(s,'Real Evolus text responses + local placement decisions · Typed-input comparison',.73,5.39,11.9,.3,13,MUTED)
text(s,'54 automated tests pass',.73,5.99,6.2,.4,22,WHITE,True)
text(s,'Next validation: speech models + public HTTPS voice demo',.73,6.47,8.8,.3,13,MUTED)
linkbox=text(s,'SOURCE REPOSITORY\ngithub.com/sg01-stack/project-zero',9.05,6.05,3.59,.62,12,PURPLE,False,PP_ALIGN.RIGHT)
for p in linkbox.text_frame.paragraphs:
    for run in p.runs:
        run.hyperlink.address='https://github.com/sg01-stack/project-zero'
        run.font.color.rgb=rgb(PURPLE)
        run.font.underline=False

# Timings are available to PowerPoint when "Use Timings" is enabled.
# Keep hyperlinks legible in viewers that apply the theme's link color.
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
for part in prs.part.package.iter_parts():
    if '/theme/' in str(part.partname) and str(part.partname).endswith('.xml'):
        theme=etree.fromstring(part.blob)
        for kind in ('hlink','folHlink'):
            for node in theme.findall(f'.//a:clrScheme/a:{kind}',ns):
                for child in list(node): node.remove(child)
                color=etree.SubElement(node,'{'+ns['a']+'}srgbClr');color.set('val',PURPLE)
        part._blob=etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
prs.save(OUT/'EchoSphere-Pitch.pptx')

lines=['# EchoSphere — six-slide presentation', '',
       '**Target: 4 minutes 10 seconds. Hard limit: 4 minutes 30 seconds.**', '',
       '[PowerPoint](presentation/EchoSphere-Pitch.pptx) · [PDF](presentation/EchoSphere-Pitch.pdf)', '',
       'The PowerPoint has editable text/vector diagrams, presenter notes and slide timings totaling 250 seconds. Use the PDF for consistent display. The spoken script below is 453 words; rehearse once and keep the final 20 seconds as a buffer. No live demo segment is included in this timing.', '',
       '| Slide | Topic | Time | Duration |', '| --- | --- | --- | --- |']
for i,c in enumerate(CONTENT,1):
    lines.append(f"| {i} | {c['title']} | {c['window']} | {c['seconds']} seconds |")
for i,c in enumerate(CONTENT,1):
    lines.extend(['',f"## {i}. {c['title']} — {c['window']}",'',c['script'],'',f"**Presenter cue (not spoken):** {c['cue']}"])
lines.extend(['','## Before recording','',
    '- Rehearse with a timer; aim to finish at 4:10. Do not add an unplanned live demo to this script.',
    '- PowerPoint: use Slide Show → Use Timings for the prepared 250-second schedule, or advance manually while following the timestamps. PDF: use full-screen mode and advance at those times.',
    '- The quiet/distant examples come from real typed Evolus tests documented in [SUBMISSION.md](SUBMISSION.md). Distances are spatial-renderer settings, not measured physical localization. Input quotes on the slides are abbreviated for readability.',
    '- The configured model label is `own/gemma-4-31b-it`. The owner reports that it runs on AMD Developer Cloud. Collect actual GPU/model-server evidence before submission; do not invent a GPU SKU, ROCm version or serving stack.',
    '- Keep the final validation line until real Whisper/Kokoro inference and the public voice demo pass. Update it only with observed results.',
    '- No public demo URL exists in this deck. The source link is real; confirm repository visibility before judging. No API keys, private endpoints or provider credentials are included.',
    '- Cost per experience and customer demand remain unmeasured. A creator/venue subscription is a proposal.',
    '- Build with Python 3 and `python-pptx`: `python3 docs/presentation/build_deck.py`. Export the resulting PowerPoint to PDF using PowerPoint or LibreOffice. These are presentation-authoring dependencies, not app runtime requirements.',''])
(OUT.parent/'PRESENTATION.md').write_text('\n'.join(lines))
print(f'Built {len(prs.slides)} slides; {sum(c["seconds"] for c in CONTENT)} seconds; '+
      f'{sum(len(c["script"].split()) for c in CONTENT)} spoken words.')
